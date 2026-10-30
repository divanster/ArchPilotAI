"""Synthetic PDF and DOCX builders for ingestion tests; no binary fixtures are checked in."""

from io import BytesIO

import docx
from pypdf import PdfReader, PdfWriter


def _pdf_string(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return b"(" + escaped.encode("latin-1") + b")"


def _text_stream(lines: list[str]) -> bytes:
    ops = [b"BT /F1 12 Tf 72 720 Td 14 TL"]
    ops += [_pdf_string(line) + b" Tj T*" for line in lines]
    ops.append(b"ET")
    return b"\n".join(ops)


def make_pdf(
    pages: list[list[str]], *, raw_streams: list[bytes] | None = None, catalog_extra: bytes = b""
) -> bytes:
    """Build a minimal PDF with one Helvetica text line per list item.

    ``raw_streams`` replaces the generated page content streams (for text-free pages);
    ``catalog_extra`` is appended to the catalog dictionary (for example an OpenAction).
    """
    streams = raw_streams if raw_streams is not None else [_text_stream(p) for p in pages]
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    pages_id = len(objects) + 2 * len(streams) + 1
    kids: list[int] = []
    for stream in streams:
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        kids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
                % (pages_id, font, content)
            )
        )
    refs = b" ".join(b"%d 0 R" % k for k in kids)
    add(b"<< /Type /Pages /Kids [" + refs + b"] /Count %d >>" % len(kids))
    catalog = add(b"<< /Type /Catalog /Pages %d 0 R " % pages_id + catalog_extra + b">>")

    out = BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, catalog, xref)
    )
    return out.getvalue()


def make_textless_pdf(pages: int = 1) -> bytes:
    """A scan-like PDF: pages carry only drawn graphics and no text layer."""
    return make_pdf([], raw_streams=[b"0 0 1 rg 72 72 300 400 re f"] * pages)


def make_encrypted_pdf(
    lines: list[str], user_password: str = "secret", algorithm: str | None = None
) -> bytes:
    writer = PdfWriter(clone_from=PdfReader(BytesIO(make_pdf([lines]))))
    if algorithm is None:
        writer.encrypt(user_password, "owner-password")
    else:
        writer.encrypt(user_password, "owner-password", algorithm=algorithm)
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(
    paragraphs: list[str], table: list[list[str]] | None = None, after: list[str] | None = None
) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    for text in after or []:
        document.add_paragraph(text)
    out = BytesIO()
    document.save(out)
    return out.getvalue()
